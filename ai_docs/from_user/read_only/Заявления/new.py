"""
Генератор документов для объяснительных записок и заявлений о пропусках занятий.

Этот модуль создает Word документы (DOCX) на основе данных о пропущенных занятиях,
загружаемых из конфигурационных файлов в форматах YAML или JSON.
"""

import json
from datetime import datetime
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt

# ============================================================================
# КОНСТАНТЫ
# ============================================================================

# Отступы документа (поля страницы)
MARGIN_TOP = Inches(0.69)
MARGIN_BOTTOM = Inches(0.5)
MARGIN_LEFT = Inches(0.89)
MARGIN_RIGHT = Inches(0.79)

# Настройки шрифта
FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(12)

# Выравнивание абзацев
ALIGN_LEFT = WD_PARAGRAPH_ALIGNMENT.LEFT
ALIGN_CENTER = WD_PARAGRAPH_ALIGNMENT.CENTER
ALIGN_RIGHT = WD_PARAGRAPH_ALIGNMENT.RIGHT
ALIGN_JUSTIFY = WD_PARAGRAPH_ALIGNMENT.DISTRIBUTE

# Размеры страницы A4
PAGE_HEIGHT_A4 = Inches(11.69)
PAGE_WIDTH_A4 = Inches(8.27)

# Межстрочный интервал
LINE_SPACING_SINGLE = 1.0

# Отступ для выравнивания подписи
SIGNATURE_SPACING = 80

# Сопоставление сокращений типов занятий с полными названиями (для объяснительной)
LESSON_TYPE_FULL_NAMES: dict[str, str] = {
    "л": "лекцию",
    "п": "практическое занятие",
    "лаб": "практическое занятие",
    "сем": "практическое занятие",
}

# Типы занятий для лекций (влияют на КНЛ)
LECTURE_TYPES: set[str] = {"л"}

# Типы занятий для семинаров/практик (влияют на КНС)
SEMINAR_TYPES: set[str] = {"п", "сем", "лаб"}

# Названия файлов конфигурации
CONFIG_FILENAME = "config"
MISSED_CLASSES_FILENAME = "missed_classes"

# ============================================================================
# ЗАГРУЗКА ДАННЫХ ИЗ ФАЙЛОВ
# ============================================================================


def load_data_file(filename_base: str, default_data: dict[str, Any]) -> dict[str, Any]:
    """
    Загружает данные из YAML или JSON файла.

    Функция пытается загрузить данные из файлов в следующем порядке приоритета:
    1. {filename_base}.yaml
    2. {filename_base}.yml
    3. {filename_base}.json

    Args:
        filename_base: Базовое имя файла без расширения
        default_data: Данные по умолчанию, используемые если файл не найден

    Returns:
        Словарь с загруженными данными или default_data при ошибке

    Example:
        >>> config = load_data_file("config", {"director": ""})
    """
    possible_extensions = [".yaml", ".yml", ".json"]

    for extension in possible_extensions:
        filepath = Path(f"{filename_base}{extension}")

        try:
            with filepath.open("r", encoding="utf-8") as file:
                if extension in (".yaml", ".yml"):
                    return yaml.safe_load(file)
                if extension == ".json":
                    return json.load(file)

        except FileNotFoundError:
            continue

        except (yaml.YAMLError, json.JSONDecodeError):
            continue

        except Exception:
            continue

    # Если ни один файл не был успешно загружен
    [f"{filename_base}{ext}" for ext in possible_extensions]
    return default_data


def load_config() -> dict[str, str]:
    """
    Загружает конфигурацию приложения из файла config.yaml/yml/json.

    Returns:
        Словарь с конфигурационными данными:
        - director: ФИО директора/помощника директора
        - student_name: ФИО студента
        - course: Курс обучения
        - group: Номер группы
        - specialty: Специальность/направление подготовки
    """
    default_config = {
        "director": "",
        "student_name": "",
        "course": "",
        "group": "",
        "specialty": "",
    }
    return load_data_file(CONFIG_FILENAME, default_config)


def load_missed_classes() -> dict[str, Any]:
    """
    Загружает данные о пропущенных занятиях из файла missed_classes.yaml/yml/json.

    Returns:
        Словарь с данными о пропусках:
        - reason: Словарь с причиной в двух падежах (instrumental, genitive)
        - document_proof: Подтверждающий документ
        - signature_date: Дата подписи (если не указана, используется текущая)
        - generate_explanatory: Флаг для генерации объяснительной (по умолчанию false)
        - disciplines: Словарь {название_дисциплины: {тип_занятия: [даты]}}
    """
    default_missed_classes = {
        "reason": {"instrumental": "", "genitive": ""},
        "document_proof": "",
        "signature_date": "",
        "generate_explanatory": False,
        "disciplines": {},
    }
    return load_data_file(MISSED_CLASSES_FILENAME, default_missed_classes)


# ============================================================================
# УТИЛИТЫ ДЛЯ РАБОТЫ С DOCX ДОКУМЕНТАМИ
# ============================================================================


def add_styled_paragraph(
    document: Document,
    text: str,
    is_bold: bool = False,
    alignment: WD_PARAGRAPH_ALIGNMENT = ALIGN_LEFT
) -> None:
    """
    Добавляет абзац с заданным стилем в документ.

    Args:
        document: Объект документа Word
        text: Текст абзаца
        is_bold: Сделать текст жирным
        alignment: Выравнивание абзаца
    """
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)

    if is_bold:
        run.bold = True

    run.font.size = FONT_SIZE
    run.font.name = FONT_NAME

    paragraph.paragraph_format.line_spacing = LINE_SPACING_SINGLE
    paragraph.alignment = alignment


def add_blank_paragraphs(document: Document, count: int = 1) -> None:
    """
    Добавляет пустые абзацы для создания вертикального отступа.

    Args:
        document: Объект документа Word
        count: Количество пустых абзацев для добавления
    """
    for _ in range(count):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = LINE_SPACING_SINGLE


def setup_document_page(document: Document) -> None:
    """
    Настраивает параметры страницы документа (размер A4, поля).

    Args:
        document: Объект документа Word для настройки
    """
    section = document.sections[0]

    # Установка размера страницы A4
    section.page_height = PAGE_HEIGHT_A4
    section.page_width = PAGE_WIDTH_A4

    # Установка полей страницы
    section.top_margin = MARGIN_TOP
    section.bottom_margin = MARGIN_BOTTOM
    section.left_margin = MARGIN_LEFT
    section.right_margin = MARGIN_RIGHT


# ============================================================================
# ОБРАБОТКА ДАННЫХ О ЗАНЯТИЯХ
# ============================================================================


def format_lesson_info_by_type(
    lesson_types_with_dates: dict[str, list[str]]
) -> tuple[dict[str, list[str]], dict[str, list[str]]]:
    """
    Разделяет занятия на лекции и практики/семинары с их датами.

    Args:
        lesson_types_with_dates: Словарь {тип_занятия: [список_дат]}

    Returns:
        Кортеж (лекции, практики) где каждый элемент - словарь {тип: [даты]}

    Example:
        >>> format_lesson_info_by_type({"л": ["01.10"], "п": ["02.10"]})
        ({"л": ["01.10"]}, {"п": ["02.10"]})
    """
    lectures = {}
    seminars = {}

    for lesson_type_abbr, dates in lesson_types_with_dates.items():
        if lesson_type_abbr.lower() in LECTURE_TYPES:
            lectures[lesson_type_abbr] = dates
        elif lesson_type_abbr.lower() in SEMINAR_TYPES:
            # Объединяем все практики/семинары/лабы под одним типом
            if "п" not in seminars:
                seminars["п"] = []
            seminars["п"].extend(dates)

    return lectures, seminars


def format_lesson_types_and_dates(
    lesson_types_with_dates: dict[str, list[str]]
) -> tuple[str, str]:
    """
    Форматирует типы занятий (без дублирования) и все даты.

    Args:
        lesson_types_with_dates: Словарь {тип_занятия: [список_дат]}

    Returns:
        Кортеж (типы_занятий_строка, даты_строка)

    Example:
        >>> format_lesson_types_and_dates({"л": ["01.10", "08.10"], "п": ["03.10"]})
        ("лекцию, практическое занятие", "01.10, 08.10, 03.10")
    """
    unique_types = []
    all_dates = []

    for lesson_type_abbr, dates in lesson_types_with_dates.items():
        # Добавляем тип занятия только один раз
        full_lesson_type = LESSON_TYPE_FULL_NAMES.get(
            lesson_type_abbr.lower(),
            lesson_type_abbr
        )
        if full_lesson_type not in unique_types:
            unique_types.append(full_lesson_type)

        # Добавляем все даты
        all_dates.extend(dates)

    formatted_types = ", ".join(unique_types)
    formatted_dates = ", ".join(all_dates)

    return formatted_types, formatted_dates


def determine_knl_kns(
    lesson_types_with_dates: dict[str, list[str]]
) -> str:
    """
    Определяет какие показатели указывать (КНЛ, КНС или оба).

    Args:
        lesson_types_with_dates: Словарь {тип_занятия: [список_дат]}

    Returns:
        Строка с показателями: "КНЛ", "КНС" или "КНЛ, КНС"
    """
    has_lectures = any(
        lesson_type.lower() in LECTURE_TYPES
        for lesson_type in lesson_types_with_dates
    )
    has_seminars = any(
        lesson_type.lower() in SEMINAR_TYPES
        for lesson_type in lesson_types_with_dates
    )

    if has_lectures and has_seminars:
        return "КНЛ, КНС"
    if has_lectures:
        return "КНЛ"
    if has_seminars:
        return "КНС"
    return "КНЛ, КНС"  # На всякий случай


# ============================================================================
# ГЕНЕРАЦИЯ СОДЕРЖИМОГО ДОКУМЕНТОВ
# ============================================================================


def add_document_header(
    document: Document,
    config: dict[str, str]
) -> None:
    """
    Добавляет стандартную шапку документа (адресат).

    Args:
        document: Объект документа Word
        config: Конфигурационные данные о студенте
    """
    add_styled_paragraph(
        document,
        "Директору института (помощнику директора)",
        alignment=ALIGN_RIGHT
    )
    add_styled_paragraph(
        document,
        config.get("director", ""),
        alignment=ALIGN_RIGHT
    )
    add_styled_paragraph(
        document,
        f"от обучающегося {config.get('course', '')} курса, "
        f"{config.get('group', '')} группы",
        alignment=ALIGN_RIGHT
    )
    add_styled_paragraph(
        document,
        "по специальности (направлению подготовки)",
        alignment=ALIGN_RIGHT
    )
    add_styled_paragraph(
        document,
        config.get("specialty", ""),
        alignment=ALIGN_RIGHT
    )
    add_styled_paragraph(
        document,
        config.get("student_name", ""),
        alignment=ALIGN_RIGHT
    )
    add_blank_paragraphs(document, 2)


def add_signature_line(document: Document, signature_date: str) -> None:
    """
    Добавляет строку с датой и подписью в конец документа.

    Args:
        document: Объект документа Word
        signature_date: Дата подписания документа
    """
    spacing = " " * (SIGNATURE_SPACING - len(signature_date))
    signature_text = f"Дата {signature_date}{spacing} Подпись обучающегося"

    add_styled_paragraph(
        document,
        signature_text,
        alignment=ALIGN_JUSTIFY
    )


def generate_explanatory_note_page(
    document: Document,
    config: dict[str, str],
    missed_data: dict[str, Any],
    discipline_name: str,
    lesson_types_with_dates: dict[str, list[str]]
) -> None:
    """
    Генерирует страницу объяснительной записки.

    Args:
        document: Объект документа Word
        config: Конфигурационные данные о студенте
        missed_data: Данные о пропусках (причина, документы)
        discipline_name: Название дисциплины
        lesson_types_with_dates: Словарь {тип_занятия: [даты]}
    """
    add_document_header(document, config)

    reason_instrumental = missed_data.get("reason", {}).get("instrumental", "")
    signature_date = missed_data.get("signature_date", "")
    if not signature_date:
        signature_date = datetime.now().strftime("%d.%m.%Y")

    add_styled_paragraph(
        document,
        "Объяснительная записка",
        is_bold=True,
        alignment=ALIGN_CENTER
    )
    add_blank_paragraphs(document)

    # Разделяем на лекции и практики
    lectures, seminars = format_lesson_info_by_type(lesson_types_with_dates)

    # Обрабатываем лекции
    if lectures:
        formatted_types, formatted_dates = format_lesson_types_and_dates(lectures)
        add_styled_paragraph(
            document,
            f"Я пропустил(а) {formatted_types} "
            f"по дисциплине {discipline_name};"
        )
        add_styled_paragraph(
            document,
            f"даты пропуска занятий: {formatted_dates};"
        )
        add_styled_paragraph(
            document,
            f"в связи с: {reason_instrumental}."
        )
        add_blank_paragraphs(document)

    # Обрабатываем практики/семинары (теперь все объединены)
    if seminars:
        formatted_types, formatted_dates = format_lesson_types_and_dates(seminars)
        add_styled_paragraph(
            document,
            f"Я пропустил(а) {formatted_types} "
            f"по дисциплине {discipline_name};"
        )
        add_styled_paragraph(
            document,
            f"даты пропуска занятий: {formatted_dates};"
        )
        add_styled_paragraph(
            document,
            f"в связи с: {reason_instrumental}."
        )
        add_blank_paragraphs(document)

    add_styled_paragraph(
        document,
        f"Подтверждающий документ прилагается: "
        f"{missed_data.get('document_proof', '')}."
    )
    add_blank_paragraphs(document)

    add_signature_line(document, signature_date)
    document.add_page_break()


def generate_application_page(
    document: Document,
    config: dict[str, str],
    missed_data: dict[str, Any],
    discipline_name: str,
    lesson_types_with_dates: dict[str, list[str]]
) -> None:
    """
    Генерирует страницу заявления.

    Args:
        document: Объект документа Word
        config: Конфигурационные данные о студенте
        missed_data: Данные о пропусках (причина, документы)
        discipline_name: Название дисциплины
        lesson_types_with_dates: Словарь {тип_занятия: [даты]}
    """
    add_document_header(document, config)

    reason_genitive = missed_data.get("reason", {}).get("genitive", "")
    signature_date = missed_data.get("signature_date", "")
    if not signature_date:
        signature_date = datetime.now().strftime("%d.%m.%Y")

    add_styled_paragraph(
        document,
        "Заявление",
        is_bold=True,
        alignment=ALIGN_CENTER
    )
    add_blank_paragraphs(document)

    # Определяем КНЛ/КНС
    knl_kns = determine_knl_kns(lesson_types_with_dates)

    # Разделяем на лекции и практики
    lectures, seminars = format_lesson_info_by_type(lesson_types_with_dates)

    add_styled_paragraph(
        document,
        f"Прошу не снижать {knl_kns} по дисциплине {discipline_name}, "
        f"в связи с пропуском занятий по причине: {reason_genitive}"
    )

    # Обрабатываем лекции
    if lectures:
        _, formatted_dates = format_lesson_types_and_dates(lectures)
        add_styled_paragraph(
            document,
            f"даты пропуска лекций: {formatted_dates};"
        )

    # Обрабатываем практики/семинары/лабы (теперь все вместе как "практические занятия")
    if seminars:
        _, formatted_dates = format_lesson_types_and_dates(seminars)
        add_styled_paragraph(
            document,
            f"даты пропуска практических занятий: {formatted_dates}."
        )

    add_blank_paragraphs(document)

    add_signature_line(document, signature_date)
    document.add_page_break()


# ============================================================================
# ОСНОВНАЯ ФУНКЦИЯ СОЗДАНИЯ ДОКУМЕНТОВ
# ============================================================================


def create_documents() -> None:
    """
    Основная функция для создания DOCX документов с объяснительными и заявлениями.

    Загружает конфигурацию и данные о пропусках, затем генерирует для каждой
    дисциплины документы согласно настройкам. Объяснительная записка создается
    только если в конфиге установлен флаг generate_explanatory: true.

    Raises:
        Exception: При ошибке сохранения документа
    """

    # Загрузка данных
    config = load_config()
    missed_data = load_missed_classes()

    disciplines = missed_data.get("disciplines", {})
    generate_explanatory = missed_data.get("generate_explanatory", False)

    # Установка даты подписи по умолчанию
    if not missed_data.get("signature_date"):
        missed_data["signature_date"] = datetime.now().strftime("%d.%m.%Y")

    if not disciplines:
        return


    # Информация о том, что будет сгенерировано
    if generate_explanatory:
        pass
    else:
        pass

    # Создание документа
    document = Document()
    setup_document_page(document)

    # Генерация страниц для каждой дисциплины
    for discipline_name, lesson_types_with_dates in disciplines.items():
        if not lesson_types_with_dates:
            continue


        # Генерируем объяснительную только если флаг установлен
        if generate_explanatory:
            generate_explanatory_note_page(
                document,
                config,
                missed_data,
                discipline_name,
                lesson_types_with_dates
            )

        # Заявление генерируем всегда
        generate_application_page(
            document,
            config,
            missed_data,
            discipline_name,
            lesson_types_with_dates
        )

    # Формирование имени выходного файла
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    student_name_sanitized = (
        config.get("student_name", "Неизвестный")
        .replace(" ", "_")
        .replace("__", "_")
    )

    if generate_explanatory:
        output_filename = (
            f"Заявление_Объяснительная_{student_name_sanitized}_{timestamp}.docx"
        )
    else:
        output_filename = (
            f"Заявление_{student_name_sanitized}_{timestamp}.docx"
        )

    # Сохранение документа
    try:
        document.save(output_filename)
        if generate_explanatory:
            pass
        else:
            pass

    except Exception:
        raise


# ============================================================================
# ТОЧКА ВХОДА
# ============================================================================


if __name__ == "__main__":
    create_documents()
